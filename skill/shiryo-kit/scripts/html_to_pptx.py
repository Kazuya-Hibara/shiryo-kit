#!/usr/bin/env python3
# Copyright (c) 2026 Carnot AI Inc. (MIT) — derived from carnot-tech/consulting-pptx-skill scripts/html_to_pptx.py
# Modifications for shiryo-kit (MIT): speaker notes from the HTML (aside.notes / data-notes), neutral document
# properties (author / lastModifiedBy empty or --author), actionable dependency errors, text boxes widened for
# list-style bullets drawn outside the text. Normally called via build.py.
"""HTMLデッキ（16:9・1 section = 1スライド）を、編集できる PPTX に変換する。

使い方（ふだんは build.py から呼ばれる）:
  python3 html_to_pptx.py deck.html                 # deck.pptx を同じフォルダに出力
  python3 html_to_pptx.py deck.html out.pptx [--author "名前"]

前提: 資料は HTML で仕上げてから変換する。
読み原稿（スピーカーノート）: 各スライドの <aside class="notes"> か data-notes 属性を PowerPoint のノートに入れる
  （契約は references/html-contract.md）。
仕組み: html_dump.mjs が Chrome で描画した要素の位置と見た目を書き出し、本スクリプトが PowerPoint の図形に組み立てる。
  文字 → テキストボックス（折り返し幅・行間・書体・色を保持）
  塗り・枠 → 四角形（角丸・多角形の clip-path は形状として再現）／片側だけの罫線 → 直線
  表 → PowerPoint の表（結合セル・セルの塗り・罫線・余白を保持）
  SVG・画像・背景画像 → 透過PNG（中の数値は編集できない）
依存: Node 22+ と Chrome（html_dump.mjs）、python-pptx。
"""
import datetime
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _deps import BrowserUnavailable, require_python_libs, require_node  # noqa: E402

require_python_libs()

from pptx_metrics import make_room_below_tables  # noqa: E402

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W, SLIDE_H = 12192000, 6858000
SANS, SERIF = "Yu Gothic", "Yu Mincho"  # Windows / Mac の PowerPoint にある和文書体
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 表スタイル「スタイルなし、表のグリッドなし」
ALIGN = {"left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


class Scale:
    def __init__(self, w_px):
        self.k = SLIDE_W / w_px  # EMU / px
        self.pt = 960.0 / w_px   # pt / px（13.333in × 72pt）

    def __call__(self, v):
        return Emu(int(round(v * self.k)))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def font_name(run):
    f = run.get("font", "")
    if run.get("serif") or re.search(r"Mincho|明朝|Serif", f, re.I):
        return SERIF
    return SANS


def strip_style(shape):
    """add_shape が付ける既定スタイル（影・テーマ色の枠）を外す"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def fill_runs(tf, paras, sc, first=True):
    for pi, p in enumerate(paras):
        para = tf.paragraphs[0] if (first and pi == 0) else tf.add_paragraph()
        para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
        lh = p.get("lineHeight")
        if lh:
            para.line_spacing = Pt(lh * sc.pt)
        para.space_before = para.space_after = Pt(0)
        bu = p.get("bullet")
        # HTML 側で「• 」を文字で打っている行も、PowerPoint では箇条書き書式にする（slide-rules §7.3）
        if not bu and p["runs"]:
            m = re.match(r"^([•・·–])\s+", p["runs"][0]["text"])
            if m:
                p = dict(p, runs=[dict(p["runs"][0], text=p["runs"][0]["text"][m.end():])] + p["runs"][1:])
                bu = {"char": m.group(1), "color": p["runs"][0].get("color"), "size": p["runs"][0]["size"]}
        if bu:  # 行頭記号は文字で打たず箇条書き書式に（slide-rules §7.3）。ぶら下げで2行目以降を本文の頭に揃える
            first_size = p["runs"][0]["size"] if p["runs"] else 12
            hang = p.get("hang") or first_size * 1.0
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(sc(hang))))
            pPr.set("indent", str(-int(sc(hang))))
            for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
                for el in pPr.findall(qn(tag)):
                    pPr.remove(el)
            if bu.get("color"):
                sf = etree.SubElement(pPr, qn("a:buClr"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", bu["color"].lstrip("#").upper())
            if bu.get("size") and first_size:
                etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(int(min(400, max(25, bu["size"] / first_size * 100)) * 1000)))
            if bu.get("auto"):
                an = etree.SubElement(pPr, qn("a:buAutoNum"))
                an.set("type", bu["auto"])
                if bu.get("start", 1) > 1:
                    an.set("startAt", str(bu["start"]))
            else:
                etree.SubElement(pPr, qn("a:buFont")).set("typeface", SANS)
                etree.SubElement(pPr, qn("a:buChar")).set("char", bu["char"])
        for r in p["runs"]:
            run = para.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(round(r["size"] * sc.pt * 2) / 2)
            f.bold, f.italic = r.get("bold", False), r.get("italic", False)
            if r.get("underline"):
                f.underline = True
            f.color.rgb = rgb(r.get("color", "#000000"))
            name = font_name(r)
            f.name = name
            rPr = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rPr, qn(tag))
                el.set("typeface", name)


def add_rect(slide, it, sc):
    x, y, w, h = sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"])
    if w <= 0 or h <= 0:
        return
    if it.get("poly"):
        pts = [(sc(it["x"] + px), sc(it["y"] + py)) for px, py in it["poly"]]
        fb = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=True)
        shp = fb.convert_to_shape()
    elif it.get("ellipse") or (it.get("radius", 0) >= min(it["w"], it["h"]) / 2 - 0.5 and abs(it["w"] - it["h"]) <= 1):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)  # border-radius 50% → 楕円・円
    elif it.get("radius", 0) > 0.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = min(0.5, it["radius"] / max(1, min(it["w"], it["h"])))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    strip_style(shp)
    if it.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(it["fill"])
        a = it.get("alpha", 1)
        if a < 0.99:
            clr = shp.fill._xPr.find(".//" + qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(a * 100000)))
    else:
        shp.fill.background()
    ln = it.get("line")
    if ln:
        shp.line.color.rgb = rgb(ln["color"])
        shp.line.width = Pt(max(0.25, ln["w"] * sc.pt))
        if ln.get("dash"):
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH if ln["dash"] == "dashed" else MSO_LINE.ROUND_DOT
    else:
        shp.line.fill.background()
    shp.text_frame.text = ""


def add_line(slide, it, sc):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sc(it["x1"]), sc(it["y1"]), sc(it["x2"]), sc(it["y2"]))
    strip_style(c)
    c.line.color.rgb = rgb(it["color"])
    c.line.width = Pt(max(0.25, it["w"] * sc.pt))
    if it.get("dash"):
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style = MSO_LINE.DASH if it["dash"] == "dashed" else MSO_LINE.ROUND_DOT


def add_text(slide, it, sc):
    # 改行がすべて明示（<br>・段落）で、自動の折り返しが無い文字は、PowerPoint でも折り返さない
    single = it.get("lines", 1) <= len(it["paras"])
    # 1行の文字は折り返さない（PowerPoint の字幅差で改行が増えるのを防ぐ）。複数行は元の折り返し幅に揃える
    # 複数行も PowerPoint の字幅差で行末の1字（句点など）が次行に落ちないよう、わずかに広げる
    w = it["w"] * (1.04 if single else 1.015) + 2
    x = it["x"]
    # ブラウザの行頭記号（list-style）は文字の箱の外の左に描かれる（hang = 0）。PowerPoint では記号のぶら下げ（1 字分、
    # fill_runs）が箱の内側に入るので、その分だけ箱を左に広げる（広げないと本文の幅が 1 字分足りずにはみ出す）
    ext = max((p["runs"][0]["size"] for p in it["paras"] if p.get("bullet") and not p.get("hang") and p["runs"]), default=0)
    x -= ext
    w += ext
    align = it["paras"][0].get("align", "left") if it["paras"] else "left"
    if single and align in ("center",):
        x -= (w - it["w"]) / 2
    elif single and align in ("right", "end"):
        x -= (w - it["w"])
    tb = slide.shapes.add_textbox(sc(x), sc(it["y"]), sc(w), sc(max(it["h"], 1)))
    if it.get("role") == "title":
        tb.name = "Title"  # check_deck.py がタイトルとして読む
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not single
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if it.get("vert"):
        tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
    fill_runs(tf, it["paras"], sc)


def cell_border(cell, side, b, sc):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"t": "a:lnT", "r": "a:lnR", "b": "a:lnB", "l": "a:lnL"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(tag))
    if b:
        ln.set("w", str(int(max(0.25, b["w"] * sc.pt) * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", b["color"].lstrip("#").upper())
        if b.get("dash"):
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash" if b["dash"] == "dashed" else "sysDot")
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    # a:lnL, lnR, lnT, lnB は塗りより前・この順（スキーマ順）
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for child in list(tcPr):
        if child.tag in [qn(o) for o in order[:order.index(tag)]]:
            idx = list(tcPr).index(child) + 1
    tcPr.insert(idx, ln)


def add_table(slide, it, sc):
    nr, nc = len(it["rows"]), len(it["cols"])
    if not nr or not nc:
        return
    gf = slide.shapes.add_table(nr, nc, sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    for i, w in enumerate(it["cols"]):
        tbl.columns[i].width = sc(w)
    for i, h in enumerate(it["rows"]):
        tbl.rows[i].height = sc(h)
    for row in tbl.rows:  # 結合で隠れるセルも含め、既定の18ptを外す
        for cell in row.cells:
            cell.text_frame.paragraphs[0]._p.get_or_add_endParaRPr().set("sz", "100")
    covered = set()
    for c in it["cells"]:
        r0, c0, rs, cs = c["r"], c["c"], c["rs"], c["cs"]
        if r0 >= nr or c0 >= nc:
            continue
        cell = tbl.cell(r0, c0)
        if rs > 1 or cs > 1:
            cell.merge(tbl.cell(min(nr, r0 + rs) - 1, min(nc, c0 + cs) - 1))
        covered.add((r0, c0))
        pt, pr, pb, pl = c["pad"]
        cell.margin_top, cell.margin_right, cell.margin_bottom, cell.margin_left = sc(pt), sc(pr), sc(pb), sc(pl)
        cell.vertical_anchor = VALIGN.get(c.get("valign"), MSO_ANCHOR.MIDDLE)
        if c.get("fill"):
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(c["fill"])
        else:
            cell.fill.background()
        for side, b in zip("trbl", c["borders"]):
            cell_border(cell, side, b, sc)
        if c.get("vert"):
            cell._tc.get_or_add_tcPr().set("vert", "eaVert")
        tf = cell.text_frame
        tf.word_wrap = True
        if c["paras"]:
            fill_runs(tf, c["paras"], sc)
        else:
            # 空セル（図形を重ねるセルを含む）は既定の18ptのままだと行が伸びるので、行の高さに収まる大きさにする
            h_pt = it["rows"][r0] * sc.pt - (pt + pb) * sc.pt
            size = max(1, min(8, h_pt * 0.7))
            endp = tf.paragraphs[0]._p.get_or_add_endParaRPr()
            endp.set("sz", str(int(size * 100)))


def set_neutral_props(prs, author="", title=""):
    """python-pptx 既定の作成者名などを残さない（作成者・最終更新者は空か --author の値）。"""
    cp = prs.core_properties
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.author = author or ""
    cp.last_modified_by = author or ""
    cp.title = title or ""
    cp.subject = cp.keywords = cp.comments = cp.category = ""
    cp.revision = 1
    cp.created = cp.modified = now
    cp.last_printed = now


def add_notes(prs, notes):
    """notes: スライド順の文字列のリスト（空文字はノートなし）。改行は段落になる。"""
    n = 0
    for slide, text in zip(prs.slides, notes):
        if text and text.strip():
            slide.notes_slide.notes_text_frame.text = text.strip()
            n += 1
    return n


def run_dump(html, td, text_only=False):
    here = Path(__file__).resolve().parent
    node = require_node()
    cmd = [node, str(here / "html_dump.mjs"), str(Path(html).resolve()), str(td)] + (["--text-only"] if text_only else [])
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", timeout=900)
    except subprocess.TimeoutExpired:
        sys.exit("html_dump.mjs が 15 分で終わりませんでした。HTML が重すぎないか（画像・SVG の数）を確かめてください。")
    if r.returncode == 3:  # ブラウザを起動・接続できない（html_dump.mjs の openPage の失敗）
        raise BrowserUnavailable(f"Chrome を起動できない: {r.stderr.strip()[-1500:]}")
    if r.returncode != 0:
        sys.exit(f"html_dump.mjs 失敗: {r.stderr.strip()[-1500:]}")
    return json.loads((Path(td) / "dump.json").read_text(encoding="utf-8"))


def assemble(data, deck, notes, out, author="", title=None, img_dir=None):
    """描画の記録（html_dump.mjs の dump.json と同じ形）から PPTX を組み立てて保存する。戻り値 (スライド数, ノートの枚数)。
    html_to_pptx_static.py（Chrome を使わない予備の書き出し）も同じ形の記録を作ってここを通す。"""
    if len(data["slides"]) != len(deck.slides):
        sys.exit(f"スライド数が合いません（描画 {len(data['slides'])} 枚 / HTML の読み取り {len(deck.slides)} 枚）。"
                 " スライドの要素は <section class=\"s\"> を入れ子にせずに並べてください。")
    sc = Scale(data["slideW"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    layout_notes = []
    for s in data["slides"]:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(s["bg"])
        items = s["items"]
        # 重なり順: 塗り・線・画像・表は HTML の順、文字はいちばん上
        for it in items:
            t = it["t"]
            if t == "rect":
                add_rect(slide, it, sc)
            elif t == "line":
                add_line(slide, it, sc)
            elif t == "img":
                p = Path(img_dir or ".") / it["file"]
                if p.exists():
                    slide.shapes.add_picture(str(p), sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
            elif t == "table":
                add_table(slide, it, sc)
        for it in items:
            if it["t"] == "text" and it["paras"]:
                add_text(slide, it, sc)
        # PowerPoint は Chrome より字幅が広く、表の行が折り返して伸びる。伸びる分を先に取り、下の物をずらす
        for msg in make_room_below_tables(list(slide.shapes), SLIDE_H):
            layout_notes.append(f"p{len(prs.slides)}: {msg}")
    n_notes = add_notes(prs, notes)
    set_neutral_props(prs, author, deck.title if title is None else title)
    prs.save(out)
    for msg in layout_notes:
        print(f"注意: {msg}", file=sys.stderr)
    return len(data["slides"]), n_notes


def convert(html, out, notes=None, author="", title=None):
    """notes=None なら HTML から読む。戻り値 (スライド数, ノートを入れた枚数)。"""
    from deck_html import parse_deck
    deck = parse_deck(Path(html))
    if notes is None:
        notes = [s.notes for s in deck.slides]
    with tempfile.TemporaryDirectory() as td:
        data = run_dump(html, td)
        return assemble(data, deck, notes, out, author, title, img_dir=td)


def main():
    args = sys.argv[1:]
    author = ""
    if "--author" in args:
        k = args.index("--author")
        author = args[k + 1]
        del args[k:k + 2]
    if not args:
        sys.exit(__doc__)
    html = args[0]
    out = args[1] if len(args) > 1 else str(Path(html).with_suffix(".pptx"))
    n, n_notes = convert(html, out, author=author)
    print(f"{n} slides（ノート {n_notes} 枚） → {out}")
    print("次: check.py pptx で読み戻して検査し、PowerPoint で開いて折り返しと重なりを確認する")


if __name__ == "__main__":
    main()
