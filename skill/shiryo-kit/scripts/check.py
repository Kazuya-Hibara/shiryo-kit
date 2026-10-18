#!/usr/bin/env python3
"""shiryo-kit: 機械チェック。どのサブコマンドも exit 0 = 合格、1 = 不合格（理由を 1 行ずつ stdout に出す）。

使い方:
  python3 check.py deck   deck.html                  # 構成・文面の検査（タイトル・プレースホルダー・避ける言い回し 等）
  python3 check.py ledger ledger.md deck.html [読み原稿.md 配布資料.docx …]   # 数字・日付・率・金額・法令名・条番号が台帳と合うか
                                                     #（旧形式 ledger deck.html ledger.md も可）
  python3 check.py pii    資料メモ/ [--allow 許可リスト.txt]   # メモリーに個人情報が残っていないか
  python3 check.py pptx   deck.pptx deck.html [--static] [--overflow warn]   # 書き出した pptx の読み戻し検査
  python3 check.py review-evidence review/            # 別の文脈で検品した記録が揃っているか（揃わなければ「未検品」）
  python3 check.py manifest out/ review/              # 渡す file が、検品した版から書き出したものか（sha256 で照合）
  ledger に --sources sources/ を付けると、確認済みの引用を保存した原典の本文と照らす（省略時は台帳の隣の sources/）

行の頭: FAIL = 直す / WARN = 目で確かめる（WARN だけなら exit 0）/ OK = 合格の要約。
  deck   に --strict を付けると WARN も不合格にする。
  pptx   は Chrome で HTML を描いて「描かれた文字」と pptx の文字を比べる（--static は Chrome を使わず HTML を静的に読む）。
         はみ出しは推定（字数 × 字の大きさ ÷ 箱の幅）。--overflow warn で WARN に下げる。
依存: deck / ledger / pii / review-evidence / manifest = 標準ライブラリだけ。pptx = python-pptx・lxml（+ 既定では Node.js 22 以上と Chrome）。
"""
import math
import re
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from deck_html import parse_deck  # noqa: E402

PLACEHOLDER = re.compile(r"(?i)\b(?:TODO|TBD|FIXME|XXX+)\b|lorem\s+ipsum|[◯○●]{2,}|△△|□□|×××|ここに[^。]{0,12}(?:入れ|書|入力|記入)|"
                         r"\?\?\?|？？？|＿{3,}|_{3,}|【(?:要記入|仮|ダミー|TBD)】|\{\{[^}]*\}\}")


# 流用元の規約のうち、この skill の判断の軸（references/doctrine.md）と合わないもの
UPSTREAM_DROP = ("タイトルがですます調",)            # です・ます のタイトルは可（そろえるかは下の混在チェックで見る）
UPSTREAM_TO_WARN = ("1つの文章塊に2文以上",)         # 配る資料・読む資料では 2 文の塊もある（文字量の規制はしない）
POLITE_END = re.compile(r"(?:です|ます|でした|ました|ません|でしょう|ましょう|ください)[。．.!！?？]?$")
TOPIC_TITLE = re.compile(r"(?:の概要|概要|の現状|現状|について|とは|の背景|背景|一覧|まとめ|のポイント|ポイント|の流れ|の仕組み)[。]?$")
TOC_TITLE = re.compile(r"^(?:目次|アジェンダ|本日の内容|本日の流れ|付録|参考資料|Appendix)", re.I)


def emit(fails, warns, oks=()):
    for f in fails:
        print(f"FAIL {f}")
    for w in warns:
        print(f"WARN {w}")
    for o in oks:
        print(f"OK {o}")
    return 1 if fails else 0


# ---------------------------------------------------------------- deck
def load_phrases(path):
    sect, out = None, {"fail": [], "warn": []}
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        m = re.fullmatch(r"\[(fail|warn)\]", line)
        if m:
            sect = m.group(1)
            continue
        if sect:
            out[sect].append(re.compile(line[3:]) if line.startswith("re:") else re.compile(re.escape(line), re.I))
    return out


def cmd_deck(path, strict=False, phrases_file=None):
    import check_deck as cd
    html = Path(path).read_text(encoding="utf-8", errors="replace")
    deck = parse_deck(html)
    fails, warns = [], []
    if not deck.slides:
        return emit(["スライドが見つからない（1 枚ごとに <section class=\"s\"> で囲む — references/html-contract.md）"], [])
    # 1) 流用元の規約チェック（HTML の版面・タイトル・本文）。読み原稿は外して検査する
    cd.fails.clear()
    cd.warns.clear()
    cd.TEMPLATE_MODE = False
    cd.PART_TYPE_NAMES = cd._part_type_names()
    titles = cd.check_html(path, html)
    cd.check_title_variety(titles)
    for f in cd.fails:
        if any(k in f for k in UPSTREAM_DROP):
            continue
        (warns if any(k in f for k in UPSTREAM_TO_WARN) else fails).append(f)
    warns += [w for w in cd.warns if "タイトルが空" not in w]
    # 2) タイトル: 表紙・章扉（.cover / .chap）以外は必須
    for s in deck.slides:
        if not s.title and not s.is_cover:
            fails.append(f"p{s.index}: タイトル（h1 か .ttl/.title/.msg）が無い（結論を 1 文で書く）")
    # 3) プレースホルダー・書きかけ（スライド本文。読み原稿は文面の規約の対象外 — 数字は check.py ledger で見る）
    for s in deck.slides:
        for m in PLACEHOLDER.finditer(s.text):
            fails.append(f"p{s.index} 本文: 書きかけ・プレースホルダーが残っている「{m.group(0)}」")
    # 2b) タイトルの形: 名詞だけ（話題の名前）・です/ます と言い切りの混在（references/doctrine.md §3・§5）
    body = [s for s in deck.slides if s.title and not s.is_cover and not TOC_TITLE.search(s.title)]
    for s in body:
        if TOPIC_TITLE.search(s.title) and len(s.title) <= 16:
            warns.append(f"p{s.index}: タイトルが話題の名前だけに見える「{s.title}」（結論を 1 文で: 「だから何か」を書く）")
    polite = [s.index for s in body if POLITE_END.search(s.title)]
    plain = [s.index for s in body if not POLITE_END.search(s.title)]
    if polite and plain:
        warns.append(f"タイトルの語尾が混在: です・ます p{polite} / 言い切り p{plain}（資料の中でどちらかにそろえる）")
    # 4) 避ける言い回し（data/banned_phrases.txt。スライド本文と読み原稿）
    ph = load_phrases(phrases_file or HERE / "data" / "banned_phrases.txt")
    for s in deck.slides:
        for where, text in (("本文", s.text), ("読み原稿", s.notes)):
            for lvl, bucket in (("fail", fails), ("warn", warns)):
                hits = sorted({m.group(0) for rx in ph[lvl] for m in rx.finditer(text)})
                if hits:
                    bucket.append(f"p{s.index} {where}: 避ける言い回し「{'」「'.join(hits[:5])}」（素直な言い方に直す）")
    # 5) 読み原稿の抜け（一部のスライドにだけある時）
    with_notes = [s.index for s in deck.slides if s.notes.strip()]
    if with_notes:
        missing = [s.index for s in deck.slides if not s.notes.strip() and not s.is_cover]
        if missing:
            warns.append(f"読み原稿が無いスライド: p{missing}（話す資料なら書く）")
    if strict:
        fails, warns = fails + warns, []
    return emit(fails, warns, [] if fails else [f"deck: {len(deck.slides)} 枚（タイトル {sum(1 for s in deck.slides if s.title)} 枚・読み原稿 {len(with_notes)} 枚）"])


# ---------------------------------------------------------------- ledger
def docx_text(path):
    """docx の本文（段落ごとに 1 行）。標準ライブラリだけで読む。"""
    import zipfile
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        p = re.sub(r"<w:(?:br|cr)\b[^>]*/>", "\n", p)
        paras.append("".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", p)))
    import html as _h
    return _h.unescape("\n".join(paras))


def cmd_ledger(args, sources=None):
    from ledger import parse_ledger, validate_rows, check_units, deck_units, text_units, load_sources, verify_quotes
    if len(args) >= 2 and args[0].lower().endswith((".html", ".htm")) and args[1].lower().endswith((".md", ".markdown", ".txt")):
        ledger_path, targets = args[1], [args[0]] + args[2:]   # 旧形式: ledger deck.html ledger.md [file…]
    else:
        ledger_path, targets = args[0], args[1:]
    if not Path(ledger_path).exists():
        return emit([f"台帳が見つからない: {ledger_path}"], [])
    rows, errors = parse_ledger(Path(ledger_path).read_text(encoding="utf-8", errors="replace"))
    fails = errors + validate_rows(rows)
    warns = []
    units, n_slides = [], 0
    for t in targets:
        p = Path(t)
        if not p.exists():
            fails.append(f"照合する file が見つからない: {t}")
            continue
        if p.resolve() == Path(ledger_path).resolve():
            continue
        suf = p.suffix.lower()
        multi = len(targets) > 1
        if suf in (".html", ".htm"):
            deck = parse_deck(p)
            n_slides += len(deck.slides)
            units += deck_units(deck, p.name if multi else "")
        elif suf == ".docx":
            units += text_units(docx_text(p), p.name)
        else:
            units += text_units(p.read_text(encoding="utf-8", errors="replace"), p.name)
    fails += check_units(units, rows)
    # 確認済みの引用を、保存した原典の本文と照らす（--sources が無ければ台帳の隣の sources/）
    n_confirmed = sum(1 for r in rows if r.status == "確認済み")
    src_paths = [Path(x) for x in sources] if sources else [Path(ledger_path).resolve().parent / "sources"]
    missing_src = [str(x) for x in src_paths if not x.exists()]
    if sources and missing_src:
        fails.append(f"--sources が見つからない: {', '.join(missing_src)}")
    src = load_sources([x for x in src_paths if x.exists()])
    if src:
        fails += verify_quotes(rows, src)
    elif n_confirmed:
        warns.append(f"確認済み {n_confirmed} 行の引用を原典の本文と照らしていない（台帳の隣に sources/ が無い。"
                     f"取得した本文を sources/raw/ に保存するか --sources で場所を渡す）")
    counts = Counter(r.status for r in rows)
    ok = [f"ledger: {len(rows)} 行（確認済み {counts.get('確認済み', 0)} / 未確認 {counts.get('未確認', 0)} / 本人提供 {counts.get('本人提供', 0)}）"
          f"・照合した file {len(targets)}（スライド {n_slides} 枚）・原典の本文 {len(src)} file"]
    return emit(fails, warns, [] if fails else ok)


# ---------------------------------------------------------------- pii
def cmd_pii(target, allow):
    from pii import scan_path, mask
    p = Path(target)
    if not p.exists():
        return emit([f"見つからない: {target}"], [])
    res = scan_path(p, allow)
    base = p if p.is_dir() else p.parent
    fails = [f"{f.relative_to(base) if f != p else f.name}:{ln}: {kind}: {mask(s)}" for f, ln, kind, s in res]
    return emit(fails, [], [] if fails else [f"pii: 個人情報らしいものは見つからなかった（{target}）"])


# ---------------------------------------------------------------- pptx
def _norm_chars(text):
    lines = []
    for line in text.splitlines():
        line = re.sub(r"^\s*[•・·–\-]\s+", "", line)  # 打った行頭記号は pptx では箇条書き書式になる
        lines.append(line)
    t = unicodedata.normalize("NFKC", "\n".join(lines))
    return re.sub(r"\s+", "", t)


def _shape_texts(shapes):
    """(テキスト, 図形) を再帰で（グループの中・表のセルも）。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(sh.shapes)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame.text, None
        elif sh.has_text_frame:
            yield sh.text_frame.text, sh


def _char_w(ch):
    if ch in " \t":
        return 0.3
    ea = unicodedata.east_asian_width(ch)
    if ea in ("W", "F", "A"):
        return 1.0
    if ch.isupper() or ch.isdigit():
        return 0.62
    return 0.52


def _para_metrics(p, default_size=18.0):
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    size = max(sizes) if sizes else default_size
    ls = p.line_spacing
    if ls is None:
        lsp = size * 1.2
    elif isinstance(ls, float):
        lsp = size * 1.2 * ls
    else:
        lsp = ls.pt
    text = "".join(r.text for r in p.runs)
    em = sum(_char_w(c) for c in text)
    marL = p._p.pPr.get("marL") if p._p.pPr is not None else None
    indent = int(marL) / 12700 if marL else 0.0
    return size, lsp, em * size, indent, text


def overflow_issues(slide, idx, slide_w, slide_h):
    out = []
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            tbl = sh.table
            total = 0.0
            for ri, row in enumerate(tbl.rows):
                need = row.height / 12700
                for ci, cell in enumerate(row.cells):
                    if cell.is_spanned:
                        continue
                    w = sum(tbl.columns[c].width for c in range(ci, ci + max(1, cell.span_width))) / 12700
                    w -= (cell.margin_left + cell.margin_right) / 12700
                    h = (cell.margin_top + cell.margin_bottom) / 12700
                    for p in cell.text_frame.paragraphs:
                        size, lsp, width, ind, text = _para_metrics(p, 10)
                        if not text:
                            continue
                        h += max(1, math.ceil(width / max(1.0, w - ind) - 0.03)) * lsp
                    if cell.span_height <= 1:
                        need = max(need, h)
                total += need
            bottom = sh.top / 12700 + total
            if bottom > slide_h / 12700 + 2:
                out.append(f"p{idx}: 表が下にはみ出しそう（推定の下端 {bottom:.0f}pt > スライド {slide_h / 12700:.0f}pt）")
            continue
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        if sh.width is None or sh.height is None:
            continue
        tf = sh.text_frame
        w = (sh.width - (tf.margin_left or 0) - (tf.margin_right or 0)) / 12700
        h = (sh.height - (tf.margin_top or 0) - (tf.margin_bottom or 0)) / 12700
        need, widest, lsmax = 0.0, 0.0, 0.0
        for p in tf.paragraphs:
            size, lsp, width, ind, text = _para_metrics(p)
            lsmax = max(lsmax, lsp)
            if not text:
                need += lsp
                continue
            avail = max(1.0, w - ind)
            widest = max(widest, width + ind)
            lines = 1 if tf.word_wrap is False else max(1, math.ceil(width / avail - 0.03))
            need += lines * lsp
        label = sh.text_frame.text.strip().replace("\n", " ")[:24]
        if tf.word_wrap is False and widest > w * 1.08 + 3:
            out.append(f"p{idx}: 1 行の文字が箱の幅を超えそう（推定 {widest:.0f}pt > 幅 {w:.0f}pt）「{label}」")
        elif need > h * 1.12 + max(3.0, lsmax * 0.35):
            out.append(f"p{idx}: 文字が箱からはみ出しそう（推定の高さ {need:.0f}pt > 箱 {h:.0f}pt）「{label}」")
        if sh.left is not None and (sh.left + sh.width > slide_w + 12700 * 4 or sh.top + sh.height > slide_h + 12700 * 4
                                    or sh.left < -12700 * 4 or sh.top < -12700 * 4):
            out.append(f"p{idx}: 文字の箱がスライドの外に出ている「{label}」")
    return out


def cmd_pptx(pptx_path, html_path, static=False, overflow="fail"):
    from _deps import require_python_libs
    require_python_libs()
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx_open import open_presentation
    prs = open_presentation(pptx_path)
    deck = parse_deck(Path(html_path))
    fails, warns = [], []
    rendered = None
    if not static:
        from html_to_pptx import run_dump
        with tempfile.TemporaryDirectory() as td:
            data = run_dump(html_path, td, text_only=True)
        rendered = [s.get("text", "") for s in data["slides"]]
    html_texts = rendered if rendered is not None else [s.text for s in deck.slides]
    slides = list(prs.slides)
    if len(slides) != len(html_texts):
        fails.append(f"スライド数が違う（pptx {len(slides)} 枚 / HTML {len(html_texts)} 枚）")
    cp = prs.core_properties
    for label, v in (("作成者", cp.author), ("最終更新者", cp.last_modified_by)):
        if v and re.search(r"python-pptx|Steve Canny", v):
            fails.append(f"文書のプロパティの{label}が既定値のまま「{v}」（build.py で書き出し直す）")
    n_pics = n_tables = n_text = 0
    for i, (slide, htext) in enumerate(zip(slides, html_texts), 1):
        ptexts = [t for t, _ in _shape_texts(slide.shapes)]
        pics = sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
        tables = sum(1 for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table)
        n_pics += pics
        n_tables += tables
        n_text += sum(1 for t in ptexts if t.strip())
        miss = Counter(_norm_chars(htext)) - Counter(_norm_chars("\n".join(ptexts)))
        if miss:
            lines = [ln.strip() for ln in htext.splitlines() if any(c in miss for c in _norm_chars(ln))]
            pic_note = f"（画像 {pics} 個: 文字が画像になっていないか）" if pics else ""
            fails.append(f"p{i}: HTML の文字が pptx に無い {sum(miss.values())} 字{pic_note}: "
                         + " / ".join(f"「{ln[:30]}」" for ln in lines[:3]))
        hs = deck.slides[i - 1] if i - 1 < len(deck.slides) else None
        slide_area = prs.slide_width * prs.slide_height
        big = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width and sh.height
               and sh.width * sh.height >= slide_area * 0.5]
        if big and (hs is None or not hs.has_media):
            warns.append(f"p{i}: スライドの半分以上を画像が占めている（HTML に img・svg が無い。CSS の背景画像でなければ、"
                         f"スライドが絵として書き出されていないか確かめる）")
        if hs is not None and hs.has_table > tables:
            fails.append(f"p{i}: HTML の表 {hs.has_table} 個に対し PowerPoint の表が {tables} 個（表が図や文字の箱になっている）")
        if hs is not None and hs.notes.strip():
            pn = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            if not pn.strip():
                fails.append(f"p{i}: HTML に読み原稿があるのに pptx のノートが空")
            elif _norm_chars(pn) != _norm_chars(hs.notes):
                fails.append(f"p{i}: pptx のノートが HTML の読み原稿と違う「{pn.strip()[:30]}」")
        issues = overflow_issues(slide, i, prs.slide_width, prs.slide_height)
        (fails if overflow == "fail" else warns).extend(issues)
    if static:
        warns.append("Chrome を使わない静的な比較（CSS で隠した文字も「欠落」に数えることがある）")
    ok = [f"pptx: {len(slides)} 枚・テキストの箱 {n_text}・表 {n_tables}・画像 {n_pics}・ノート "
          f"{sum(1 for s in slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())} 枚"]
    return emit(fails, warns, [] if fails else ok)


# ---------------------------------------------------------------- review-evidence / manifest
def _evidence(review_dir):
    """(最後の巡, FAIL 理由, WARN)。FAIL 理由が空なら、最後の巡は別の文脈の検品として認める。"""
    from review_evidence import check_review_dir
    last, earlier, errs = check_review_dir(review_dir)
    warns = []
    for r in earlier:
        if not r.ok:
            warns.append(f"前の巡（{r.n} 巡目）の記録に不備: {r.fails[0]}")
    if errs:
        return None, errs, warns
    warns += last.warns
    return last, last.fails, warns


def cmd_review_evidence(review_dir):
    last, fails, warns = _evidence(review_dir)
    if fails:
        return emit([f"未検品: {f}" for f in fails], warns)
    rec = last.record
    return emit([], warns, [f"review-evidence: {last.n} 巡目 結論={last.verdict}（{rec.get('method')}・reviewer {rec.get('reviewer_id')}・"
                            f"{rec.get('started_at')} 〜 {rec.get('ended_at')}）"])


def cmd_manifest(out_dir, review_dir):
    import json
    from review_evidence import sha256
    out = Path(out_dir)
    mpath = out / "manifest.json"
    if not mpath.exists():
        return emit([f"未検品: 書き出しの記録 {mpath} が無い（build.py で書き出し直す）"], [])
    try:
        m = json.loads(mpath.read_text(encoding="utf-8"))
    except ValueError as e:
        return emit([f"未検品: {mpath} が JSON として読めない（{e}）"], [])
    fails, warns = [], []
    deck = m.get("deck") or {}
    deck_sha = deck.get("sha256")
    if not deck_sha:
        fails.append("書き出しの記録に元の deck の sha256 が無い（build.py で書き出し直す）")
    else:
        dp = (out / deck.get("path", "")).resolve()
        if not dp.exists():
            fails.append(f"書き出しの元の {deck.get('path')} が見つからない")
        elif sha256(dp) != deck_sha:
            fails.append(f"書き出した後に {dp.name} が変わった（build.py で書き出し直してから検品し直す）")
    listed = set()
    for o in m.get("outputs", []):
        f = out / o.get("file", "")
        listed.add(f.name)
        if not f.exists():
            fails.append(f"記録にある {o.get('file')} が無い")
        elif sha256(f) != o.get("sha256"):
            fails.append(f"{f.name} が書き出した後に変わった（手で直した file は渡さない。HTML を直して書き出し直す）")
        if deck_sha and o.get("deck_sha256") != deck_sha:
            fails.append(f"{o.get('file')} は今の deck より前の版から書き出したもの")
    if not m.get("outputs"):
        fails.append("書き出した file の記録が無い")
    for f in sorted(out.iterdir()) if out.is_dir() else []:
        if f.is_file() and f.name not in listed and f.name != "manifest.json" and f.suffix.lower() in (".pptx", ".pdf"):
            fails.append(f"{f.name} は build.py の書き出しの記録に無い pptx/pdf（即席の変換器などで作った物は渡さない。"
                         "消すか out/ の外に移す。build.py で書き出し直す）")
        elif f.is_file() and f.name not in listed and f.name != "manifest.json" and f.suffix.lower() in (".md", ".docx", ".html"):
            warns.append(f"{f.name} は書き出しの記録に無い（照合できないので渡さない）")
    last, efails, ewarns = _evidence(review_dir)
    warns += ewarns
    fails += [f"未検品: {f}" for f in efails]
    if last is not None and not efails:
        h = last.record.get("inputs_sha256") or {}
        if deck_sha and h.get("deck") != deck_sha:
            fails.append(f"未検品: 書き出した deck は {last.n} 巡目で検品した deck と違う（sha256 が違う）")
        led = m.get("ledger") or {}
        if led.get("sha256") and h.get("ledger") != led["sha256"]:
            fails.append(f"未検品: 書き出した時の台帳は {last.n} 巡目で検品した台帳と違う")
        br = m.get("brief") or {}
        if br.get("sha256") and h.get("brief") != br["sha256"]:
            fails.append(f"未検品: 書き出した時の BRIEF は {last.n} 巡目で検品した BRIEF と違う")
    if fails:
        return emit(fails, warns)
    state = "検品済み" if last.verdict == "合格" else "検品済み・指摘が残っています"
    if last.verdict != "合格":
        warns.append(f"{last.n} 巡目の結論が「{last.verdict}」（直してもう一度検品するか、本人の判断を BRIEF の決定ログに残してから渡す）")
    files = ", ".join(sorted(listed))
    return emit([], warns, [f"manifest: 状態={state}（review-{last.n}・渡せる file: {files}）"])


def main():
    args = sys.argv[1:]
    if not args or args[0] in ("-h", "--help"):
        print(__doc__)
        return 0 if args else 2

    def flag(name):
        if name in args:
            args.remove(name)
            return True
        return False

    def opt(name, default=None):
        if name in args:
            k = args.index(name)
            v = args[k + 1]
            del args[k:k + 2]
            return v
        return default

    cmd = args.pop(0)
    if cmd == "deck" and args:
        strict = flag("--strict")
        phrases = opt("--phrases")
        return cmd_deck(args[0], strict, phrases)
    if cmd == "ledger":
        sources = []
        while "--sources" in args:
            sources.append(opt("--sources"))
        if len(args) >= 2:
            return cmd_ledger(args, sources)
    if cmd == "pii" and args:
        allow = []
        while "--allow" in args:
            allow.append(opt("--allow"))
        return cmd_pii(args[0], allow)
    if cmd == "pptx" and len(args) >= 2:
        static = flag("--static")
        overflow = opt("--overflow", "fail")
        return cmd_pptx(args[0], args[1], static, overflow)
    if cmd == "review-evidence" and args:
        return cmd_review_evidence(args[0])
    if cmd == "manifest" and len(args) >= 2:
        return cmd_manifest(args[0], args[1])
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())
