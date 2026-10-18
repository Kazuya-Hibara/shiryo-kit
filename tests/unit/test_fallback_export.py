"""F3: 表の伸びと下の物の重なり（CC2-1）・Chrome を使わない予備の書き出し（D4-4）・未検品の印（CC2-3）・配布物の記録（CC1-10）。"""
import json
import shutil
import unittest

from _helpers import FIXTURES, HAS_CHROME, HAS_PPTX, ROOT, SCRIPTS, TRAPS, TempDir, check, fails, run, warns

CLAIMS = FIXTURES / "claims_deck.html"  # e2e CC2 の claims の deck（架空の事務所・架空の連絡先だけ）
KATA = SCRIPTS.parent / "assets" / "kata"
FIXTURE_DECKS = [CLAIMS, TRAPS / "deck.html"] + [KATA / k / "sample.html" for k in ("seminar", "advisory", "proposal", "report")]
PARTS = SCRIPTS.parent / "assets" / "parts"
FIXTURE_DECKS += [PARTS / "freeform_parts_16x9.html", PARTS / "freeform_parts_more_16x9.html"]  # 型の外で組む時の部品集も
HAS_SOFFICE = False
try:
    import sys
    sys.path.insert(0, str(SCRIPTS))
    from office_pdf import find_soffice
    HAS_SOFFICE = bool(find_soffice())
except Exception:  # noqa: BLE001
    pass


class Metrics(unittest.TestCase):
    def test_autospace_and_line_end_waste(self):
        from pptx_metrics import wrap_lines
        # CC2 p3: Chrome では 1 行、LibreOffice では「撤/廃」で 2 行（和文と数字の間の空きで幅が足りない）
        self.assertEqual(wrap_lines("この要件は2026年10月に撤廃", 15, 209.9), 2)
        # CC2 p10: 全体の幅だけなら 2 行に収まるが、行末の余りで 3 行になる
        self.assertEqual(wrap_lines("実際の労働時間が2か月連続で週20時間以上になり、その後も続く見込みなら、3か月目から加入対象", 15, 354.9), 3)
        # 「〜」の前後には空きを入れない（p4 の誤検知を防ぐ）
        self.assertEqual(wrap_lines("36〜50人", 24, 120), 1)
        # 句読点は行頭に送らない（前の行にぶら下げる）
        self.assertEqual(wrap_lines("あいうえお。", 10, 50), 1)


def table_pptx(path, last_row_text, row_h_pt, src_gap_pt):
    """3 行の表と、そのすぐ下の出典の箱。"""
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    gf = s.shapes.add_table(3, 2, Pt(45), Pt(150), Pt(491), Pt(row_h_pt * 3))
    t = gf.table
    t.columns[0].width, t.columns[1].width = Pt(122), Pt(369)
    for i, (a, b) in enumerate((("基本", "就業規則で決まっている週の時間"), ("週ごと", "1週間の平均で見る"), ("実際", last_row_text))):
        t.rows[i].height = Pt(row_h_pt)
        for cell, text in ((t.cell(i, 0), a), (t.cell(i, 1), b)):
            cell.text = text
            cell.margin_left = cell.margin_right = Pt(7)
            cell.margin_top = cell.margin_bottom = Pt(12.8)
            for p in cell.text_frame.paragraphs:
                p.line_spacing = Pt(24.75)
                for r in p.runs:
                    r.font.size = Pt(15)
    tb = s.shapes.add_textbox(Pt(45), Pt(150 + row_h_pt * 3 + src_gap_pt), Pt(800), Pt(12))
    tb.text_frame.text = "出典: 厚生労働省 適用拡大Q&A"
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(7.5)
    prs.save(path)


@unittest.skipUnless(HAS_PPTX, "python-pptx と lxml が要る")
class TableOverlap(unittest.TestCase):
    LONG = "実際の労働時間が2か月連続で週20時間以上になり、その後も続く見込みなら、3か月目から加入対象"

    def _check(self, d):
        (d / "deck.html").write_text("<html><body><section class='s'><table><tr><td>基本</td></tr></table></section></body></html>",
                                     encoding="utf-8")
        return check("pptx", d / "t.pptx", d / "deck.html", "--static")

    def test_table_growing_into_source_line_fails(self):
        with TempDir() as d:
            table_pptx(d / "t.pptx", self.LONG, 75.4, 14)  # 最後の行は 3 行要る（宣言は 2 行分）→ 下の出典にかぶる
            code, out, _ = self._check(d)
            self.assertEqual(code, 1, out)
            self.assertTrue(any("重なりそう" in f and "出典" in f for f in fails(out)), out)

    def test_table_with_room_passes(self):
        with TempDir() as d:
            table_pptx(d / "t.pptx", "加入の対象", 75.4, 14)
            code, out, _ = self._check(d)
            self.assertFalse(any("重なりそう" in f for f in fails(out)), out)

    def test_make_room_moves_source_below_table(self):
        from pptx import Presentation
        from pptx_metrics import make_room_below_tables, table_growth
        with TempDir() as d:
            table_pptx(d / "t.pptx", self.LONG, 75.4, 14)
            prs = Presentation(d / "t.pptx")
            shapes = list(prs.slides[0].shapes)
            notes = make_room_below_tables(shapes, prs.slide_height)
            self.assertEqual(notes, [])
            tbl = next(s for s in shapes if s.has_table)
            src = next(s for s in shapes if s.has_text_frame and "出典" in s.text_frame.text)
            declared, total, _ = table_growth(tbl)
            self.assertAlmostEqual(declared, total, delta=0.5)  # 行の高さを推定に合わせた
            self.assertGreater(src.top / 12700, tbl.top / 12700 + total)

    def test_footer_is_not_moved(self):
        from pptx import Presentation
        from pptx.util import Pt
        from pptx_metrics import make_room_below_tables
        with TempDir() as d:
            table_pptx(d / "t.pptx", self.LONG * 4, 75.4, 14)  # 表が下の帯まで伸びる
            prs = Presentation(d / "t.pptx")
            s = prs.slides[0]
            foot = s.shapes.add_textbox(Pt(45), Pt(515), Pt(300), Pt(10))
            foot.text_frame.text = "フッター"
            notes = make_room_below_tables(list(s.shapes), prs.slide_height)
            self.assertEqual(round(foot.top / 12700), 515)
            self.assertTrue(notes and "重なる" in notes[0], notes)


@unittest.skipUnless(HAS_PPTX and HAS_CHROME, "python-pptx と Chrome が要る")
class ClaimsRegression(unittest.TestCase):
    def test_claims_deck_builds_without_table_overlap(self):
        with TempDir() as d:
            shutil.copy(CLAIMS, d / "deck.html")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx", "--engine", "chrome", cwd=d)
            self.assertEqual(code, 0, out + err)
            code, out, _ = check("pptx", d / "out" / "deck.pptx", d / "deck.html")
            self.assertEqual(code, 0, out)


@unittest.skipUnless(HAS_PPTX, "python-pptx と lxml が要る")
class Fallback(unittest.TestCase):
    def _readback(self, pptx, html):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from deck_html import parse_deck
        prs = Presentation(pptx)
        deck = parse_deck(html)
        self.assertEqual(len(prs.slides), len(deck.slides))
        for slide, hs in zip(prs.slides, deck.slides):
            tables = sum(1 for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table)
            pics = sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
            self.assertEqual(tables, hs.has_table, f"p{hs.index}: 表は PowerPoint の表のまま")
            self.assertEqual(pics, 0, f"p{hs.index}: 文字を画像にしていない")
            if hs.notes.strip():
                self.assertTrue(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip(), f"p{hs.index}: ノート")

    def test_fixture_decks_keep_text_tables_notes(self):
        for html in FIXTURE_DECKS:
            with self.subTest(deck=html.parent.name + "/" + html.name), TempDir() as d:
                code, out, err = run("html_to_pptx_static.py", html, d / "x.pptx")
                self.assertEqual(code, 0, out + err)
                self._readback(d / "x.pptx", html)
                code, out, _ = check("pptx", d / "x.pptx", html, "--static")
                self.assertFalse([f for f in fails(out) if "HTML の文字が pptx に無い" in f or "ノート" in f], out)
                if HAS_CHROME:  # Chrome で描いた文字と比べても欠落 0
                    code, out, _ = check("pptx", d / "x.pptx", html)
                    self.assertFalse([f for f in fails(out) if "HTML の文字が pptx に無い" in f], out)

    def test_build_records_fallback_mode(self):
        with TempDir() as d:
            shutil.copy(CLAIMS, d / "deck.html")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx,notes", "--engine", "python", cwd=d)
            self.assertEqual(code, 0, out + err)
            self.assertIn("描画の確認ができていない書き出しです", out)
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(m["export_mode"], "python-fallback")
            self.assertFalse(m["render_checked"])
            self.assertEqual({o["file"]: o["mode"] for o in m["outputs"]},
                             {"deck.pptx": "python-fallback", "deck_読み原稿.md": "text"})

    def test_auto_falls_back_without_node(self):
        import os
        with TempDir() as d:
            shutil.copy(TRAPS / "deck.html", d / "deck.html")
            env = {"PATH": os.pathsep.join(p for p in os.environ.get("PATH", "").split(os.pathsep)
                                           if not (shutil.which("node", path=p)))}
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx", env=env, cwd=d)
            self.assertEqual(code, 0, out + err)
            self.assertIn("予備に切り替えた理由", out)
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(m["export_mode"], "python-fallback")
            self.assertIn("Node.js", m.get("fallback_reason", ""))

    @unittest.skipUnless(HAS_SOFFICE, "LibreOffice が要る")
    def test_fallback_pdf_via_libreoffice_and_timeout(self):
        with TempDir() as d:
            shutil.copy(TRAPS / "deck.html", d / "deck.html")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pdf", "--engine", "python", cwd=d)
            self.assertEqual(code, 0, out + err)
            self.assertTrue((d / "out" / "deck.pdf").stat().st_size > 1000)
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual([o["mode"] for o in m["outputs"]], ["libreoffice"])
            code, out, err = run("build.py", "deck.html", "--out", "out2", "--formats", "pdf", "--engine", "python",
                                 cwd=d, env={"SHIRYO_SOFFICE_TIMEOUT": "1"})
            self.assertEqual(code, 0, out + err)
            self.assertIn("PDF は作れませんでした", out)
            self.assertFalse((d / "out2" / "deck.pdf").exists())


@unittest.skipUnless(HAS_PPTX, "python-pptx と lxml が要る")
class Unreviewed(unittest.TestCase):
    def test_mark_is_inside_the_file_and_no_unmarked_copy(self):
        from pptx import Presentation
        with TempDir() as d:
            shutil.copy(TRAPS / "deck.html", d / "deck.html")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx,notes", "--engine", "python",
                                 "--unreviewed", cwd=d)
            self.assertEqual(code, 0, out + err)
            names = sorted(p.name for p in (d / "out").iterdir())
            self.assertEqual(names, ["deck_未検品.pptx", "deck_読み原稿_未検品.md", "manifest.json"])
            prs = Presentation(d / "out" / "deck_未検品.pptx")
            for slide in prs.slides:
                self.assertTrue(any(sh.has_text_frame and sh.text_frame.text == "未検品" for sh in slide.shapes))
            self.assertIn("未検品", prs.core_properties.title)
            self.assertEqual(prs.core_properties.keywords, "未検品")
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertTrue(m.get("unreviewed_mark"))
            o = next(o for o in m["outputs"] if o["file"] == "deck_未検品.pptx")
            import hashlib
            self.assertEqual(o["sha256"], hashlib.sha256((d / "out" / "deck_未検品.pptx").read_bytes()).hexdigest())


class Handout(unittest.TestCase):
    def test_handout_is_recorded_and_needs_review(self):
        with TempDir() as d:
            for n in ("BRIEF.md", "ledger.md", "deck.html"):
                (d / n).write_bytes((TRAPS / n).read_bytes())
            (d / "handout.html").write_text(
                "<html><head><title>配布</title></head><body><section class='s'><h1>配布 1 枚</h1><p>手順は3つ</p></section></body></html>",
                encoding="utf-8")
            self.assertEqual(run("build.py", "handout.html", "--out", "out", "--formats", "notes", "--handout", cwd=d)[0], 1)
            self.assertEqual(run("build.py", "deck.html", "--out", "out", "--formats", "notes", cwd=d)[0], 0)
            code, out, err = run("build.py", "handout.html", "--out", "out", "--formats", "notes", "--handout", cwd=d)
            self.assertEqual(code, 0, out + err)
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(m["deck"]["path"], "../deck.html")  # 本体の deck はそのまま
            files = {o["file"]: o["source"] for o in m["outputs"]}
            self.assertEqual(files, {"deck_読み原稿.md": "../deck.html", "handout_読み原稿.md": "../handout.html"})
            # 配布物を検品に入れていない → 未検品
            from test_review_manifest import filled_request, reviewer_cmd, sha
            (d / "review").mkdir()
            (d / "review" / "request-1.md").write_text(filled_request(1), encoding="utf-8")
            code, out, err = run("review_record.py", "run", "review", "--round", "1", "--stdin", "review/request-1.md", "--",
                                 *reviewer_cmd(deck_sha=sha(d / "deck.html")), cwd=d)
            self.assertEqual(code, 0, out + err)
            code, out, _ = check("manifest", "out", "review", cwd=d)
            self.assertEqual(code, 1, out)
            self.assertTrue(any("配布物 handout.html は 1 巡目の検品に入っていない" in f for f in fails(out)), out)
            # --extra で一緒に検品する（依頼文にも配布物の行）→ 通る
            shutil.rmtree(d / "review")
            (d / "review").mkdir()
            (d / "review" / "request-1.md").write_text(
                filled_request(1) + "- 配布物: /work/資料/case/handout.html\n", encoding="utf-8")
            code, out, err = run("review_record.py", "run", "review", "--round", "1", "--stdin", "review/request-1.md",
                                 "--extra", "handout.html", "--", *reviewer_cmd(deck_sha=sha(d / "deck.html")), cwd=d)
            self.assertEqual(code, 0, out + err)
            code, out, _ = check("manifest", "out", "review", cwd=d)
            self.assertEqual(code, 0, out)
            self.assertIn("handout_読み原稿.md", out)
            # 配布物を直したら、書き出し直すまで渡せない
            (d / "handout.html").write_text((d / "handout.html").read_text(encoding="utf-8").replace("3つ", "4つ"), encoding="utf-8")
            code, out, _ = check("manifest", "out", "review", cwd=d)
            self.assertEqual(code, 1)
            self.assertTrue(any("handout_読み原稿.md は今の handout.html より前の版" in f for f in fails(out)), out)


ROWSPAN = ("<html><head><style>.s{width:338.67mm;height:190.5mm;position:relative;padding:10mm 16mm;box-sizing:border-box}"
           "td{padding:2mm 3mm;font-size:12pt;line-height:1.6;border-bottom:1px solid #999}.src{font-size:8pt}</style></head><body>"
           "<section class='s'><h1>結合セルの表</h1><table>"
           "<tr><td rowspan='3' style='width:120mm'>" + "就業規則の見直しは、労働基準法・労働契約法・育児介護休業法の改正に合わせて行う必要があり、届出の手続きも含めて確認する" * 2
           + "</td><td>行0</td></tr><tr><td>行1</td></tr><tr><td>行2</td></tr></table>"
           "<div class='src'>出典：rowspan の下の行</div></section></body></html>")


@unittest.skipUnless(HAS_PPTX, "python-pptx と lxml が要る")
class ReviewFixes(unittest.TestCase):
    def test_rowspan_table_does_not_cover_source(self):
        # レビュー M-2: 縦に結合したセルは、最後の行だけで中身が入る高さにする（LibreOffice の伸び方）
        from pptx import Presentation
        from pptx_metrics import table_growth
        with TempDir() as d:
            (d / "deck.html").write_text(ROWSPAN, encoding="utf-8")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx", "--engine", "python", cwd=d)
            self.assertEqual(code, 0, out + err)
            code, out, _ = check("pptx", d / "out" / "deck.pptx", d / "deck.html", "--static")
            self.assertEqual(code, 0, out)
            shapes = list(Presentation(d / "out" / "deck.pptx").slides[0].shapes)
            tbl = next(x for x in shapes if x.has_table)
            src = next(x for x in shapes if x.has_text_frame and "出典" in x.text_frame.text)
            _, total, rows = table_growth(tbl)
            self.assertGreater(rows[-1], rows[0] * 2)
            self.assertGreaterEqual(src.top / 12700, tbl.top / 12700 + total)

    def test_notes_only_build_is_text_mode(self):
        with TempDir() as d:
            shutil.copy(TRAPS / "deck.html", d / "deck.html")
            self.assertEqual(run("build.py", "deck.html", "--out", "out", "--formats", "notes", cwd=d)[0], 0)
            m = json.loads((d / "out" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual((m["export_mode"], m["render_checked"]), ("text", False))

    def test_dropped_image_is_reported(self):
        with TempDir() as d:
            (d / "deck.html").write_text("<html><body><section class='s'><h1>ロゴ</h1><img src='data:image/png;base64,AAAA'>"
                                         "<svg><text>図の字</text></svg></section></body></html>", encoding="utf-8")
            code, out, err = run("build.py", "deck.html", "--out", "out", "--formats", "pptx", "--engine", "python", cwd=d)
            self.assertEqual(code, 0, out + err)
            self.assertIn("画像・図が入っていない", out)
            code, out, _ = check("pptx", d / "out" / "deck.pptx", d / "deck.html", "--static")
            self.assertTrue(any("画像・図が 2 個" in w for w in warns(out)), out)


class StaticCss(unittest.TestCase):
    """deck_html.py（静的な読み取り）が <style> の display・visibility を後の規則まで含めて読む（B の指摘）。"""

    def test_cascade_and_override(self):
        from deck_html import parse_deck
        h = ("<html><head><style>.logo{visibility:hidden}.cover .logo{visibility:visible}.slide{display:none}"
             ".slide.active{display:block}.x{display:none}.x .y{display:block}</style></head><body>"
             "<section class='s cover'><div class='logo'>表紙の題</div><p class='x'>消える<span class='y'>これも消える</span></p></section>"
             "<section class='s'><div class='logo'>隠れる</div><h1>本文</h1></section>"
             "<div class='slide'><h1>スライド2</h1></div></body></html>")
        self.assertEqual([s.text for s in parse_deck(h).slides], ["表紙の題", "本文", "スライド2"])

    def test_skip_marker_reaches_reading_script_md(self):
        from build import notes_markdown
        from deck_html import parse_deck
        d = parse_deck("<html><body><section class='s'><h1>題</h1><aside class='notes'><p>本当の話です。</p>"
                       "<p data-ledger='skip'>たとえば架空の会社で月給30万円の人なら</p></aside></section></body></html>")
        md = notes_markdown(d, [s.notes for s in d.slides])
        self.assertIn("月給30万円の人なら <!-- ledger:skip -->", md)
        self.assertIn("本当の話です。\n", md)


class Versions(unittest.TestCase):
    HEAD = "| 版 | 日時 | 変わった所 | きっかけ | deck.html の sha256 | 書き出した file (sha256) | 検品 | 状態 |\n|---|---|---|---|---|---|---|---|\n"

    def _run(self, rows):
        with TempDir() as d:
            (d / "versions.md").write_text("# 版の記録\n\n" + self.HEAD + rows, encoding="utf-8")
            return check("versions", d / "versions.md")

    def test_good_rows_pass(self):
        code, out, _ = self._run("| v1 | 10/18 21:10 | 初稿 | — | 3fa1… | out/deck.pptx (91c2…) | review-1: 直しが必要 | 作業中 |\n"
                                 "| v2 | 10/18 21:40 | 3 枚目 \\| 5 枚目 | review-1 | 77b0… | out/deck.pptx (a4d9…) | review-2: 合格 | 検品済み |\n")
        self.assertEqual(code, 0, out)
        self.assertIn("2 行・8 列", out)

    def test_shifted_row_fails(self):
        # CC2-4: v3 の行で「本人の指示」と sha が別の列に入り、セルが 1 つ多い
        code, out, _ = self._run("| v3 | 10/19 | 3 枚目 | 本人の指示 | 9ccc… | out/deck.pptx (a118…) | review-3: 合格 | 検品済み | 納品 |\n")
        self.assertEqual(code, 1, out)
        self.assertTrue(any("列が 9 個（見出しは 8 個）" in f for f in fails(out)), out)

    def test_unknown_state_warns(self):
        code, out, _ = self._run("| v1 | 10/18 | 初稿 | — | 3fa1… | out/deck.pptx | review-1 | 差し替え済み |\n")
        self.assertEqual(code, 0, out)
        self.assertTrue(warns(out), out)


if __name__ == "__main__":
    unittest.main()
