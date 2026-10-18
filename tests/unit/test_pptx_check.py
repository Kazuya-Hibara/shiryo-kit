"""check.py pptx: python-pptx で作った pptx を HTML と比べる（--static）。欠落・絵だけの枚・表・ノート・はみ出し・文書のプロパティ。
Chrome がある時は build.py で書き出した本物の pptx でも通し、壊した版が落ちることを確かめる。"""
import struct
import unittest
import zlib

from _helpers import FIXTURES, HAS_CHROME, HAS_PPTX, TempDir, check, deck_html, fails, run, warns

S1 = "<h1>就業規則は年1回見直します</h1><p>改正のたびに一行足すだけで防げます</p>"
S1_NOTES = S1 + '<aside class="notes"><p>今日の結論です。</p></aside>'
TABLE = "<h1>手続きは3つで終わります</h1><table><tr><th>手続き</th><th>期限</th></tr><tr><td>届出</td><td>10日以内</td></tr></table>"


def png_bytes(w=8, h=8):
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def make_pptx(path, slides, author="", png=None):
    """slides: [{"boxes": [(text, w_in, h_in, size_pt)], "table": [[...]], "picture": True, "notes": "..."}]"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    if author is not None:
        prs.core_properties.author = author
        prs.core_properties.last_modified_by = author
    for spec in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        top = 0.5
        for text, w, h, size in spec.get("boxes", []):
            tb = s.shapes.add_textbox(Inches(0.5), Inches(top), Inches(w), Inches(h))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.text = text
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
            top += h + 0.2
        if spec.get("table"):
            rows = spec["table"]
            t = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(top), Inches(8), Inches(0.5 * len(rows))).table
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
        if spec.get("picture"):
            s.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(path)


@unittest.skipUnless(HAS_PPTX, "python-pptx と lxml が要る")
class StaticPptx(unittest.TestCase):
    def run_case(self, html_slides, pptx_slides, args=(), author=""):
        with TempDir() as d:
            (d / "deck.html").write_text(deck_html(*html_slides), encoding="utf-8")
            png = d / "x.png"
            png.write_bytes(png_bytes())
            make_pptx(d / "deck.pptx", pptx_slides, author=author, png=str(png))
            return check("pptx", d / "deck.pptx", d / "deck.html", "--static", *args)

    GOOD = {"boxes": [("就業規則は年1回見直します", 10, 0.8, 24), ("改正のたびに一行足すだけで防げます", 10, 0.8, 18)],
            "notes": "今日の結論です。"}

    def test_matching_pptx_passes(self):
        code, out, _ = self.run_case([S1_NOTES], [self.GOOD])
        self.assertEqual(code, 0, out)
        self.assertIn("OK pptx: 1 枚", out)

    def test_dropped_text_fails(self):
        spec = {"boxes": [("就業規則は年1回見直します", 10, 0.8, 24)], "notes": "今日の結論です。"}
        code, out, _ = self.run_case([S1_NOTES], [spec])
        self.assertEqual(code, 1)
        self.assertTrue(any("HTML の文字が pptx に無い" in f for f in fails(out)), out)

    def test_picture_only_slide_fails(self):
        code, out, _ = self.run_case([S1_NOTES], [{"picture": True, "notes": "今日の結論です。"}])
        self.assertEqual(code, 1)
        self.assertTrue(any("pptx に無い" in f and "画像 1 個" in f for f in fails(out)), out)
        self.assertTrue(any("半分以上を画像" in w for w in warns(out)), out)

    def test_table_as_text_boxes_fails(self):
        spec = {"boxes": [("手続きは3つで終わります", 10, 0.8, 24), ("手続き 期限 届出 10日以内", 10, 0.8, 18)]}
        code, out, _ = self.run_case([TABLE], [spec])
        self.assertEqual(code, 1)
        self.assertTrue(any("PowerPoint の表が 0 個" in f for f in fails(out)), out)

    def test_native_table_passes(self):
        spec = {"boxes": [("手続きは3つで終わります", 10, 0.8, 24)], "table": [["手続き", "期限"], ["届出", "10日以内"]]}
        code, out, _ = self.run_case([TABLE], [spec])
        self.assertEqual(code, 0, out)

    def test_missing_notes_fails(self):
        spec = dict(self.GOOD, notes="")
        code, out, _ = self.run_case([S1_NOTES], [spec])
        self.assertEqual(code, 1)
        self.assertTrue(any("ノートが空" in f for f in fails(out)), out)

    def test_different_notes_fails(self):
        spec = dict(self.GOOD, notes="別の原稿です。")
        code, out, _ = self.run_case([S1_NOTES], [spec])
        self.assertTrue(any("ノートが HTML の読み原稿と違う" in f for f in fails(out)), out)

    def test_overflowing_textbox_fails(self):
        long = "改正のたびに一行足すだけで防げます" * 6
        spec = {"boxes": [("就業規則は年1回見直します", 10, 0.8, 24), (long, 3, 0.5, 18)], "notes": "今日の結論です。"}
        html = f"<h1>就業規則は年1回見直します</h1><p>{long}</p><aside class=\"notes\"><p>今日の結論です。</p></aside>"
        code, out, _ = self.run_case([html], [spec])
        self.assertEqual(code, 1)
        self.assertTrue(any("はみ出しそう" in f for f in fails(out)), out)
        code, out, _ = self.run_case([html], [spec], args=("--overflow", "warn"))
        self.assertEqual(code, 0, out)
        self.assertTrue(any("はみ出しそう" in w for w in warns(out)), out)

    def test_box_outside_slide_fails(self):
        from pptx import Presentation
        from pptx.util import Inches
        with TempDir() as d:
            (d / "deck.html").write_text(deck_html(S1), encoding="utf-8")
            make_pptx(d / "deck.pptx", [{"boxes": [("就業規則は年1回見直します", 10, 0.8, 24), ("改正のたびに一行足すだけで防げます", 10, 0.8, 18)]}])
            prs = Presentation(d / "deck.pptx")
            prs.slides[0].shapes[1].left = Inches(12)
            prs.save(d / "deck.pptx")
            code, out, _ = check("pptx", d / "deck.pptx", d / "deck.html", "--static")
        self.assertTrue(any("スライドの外" in f for f in fails(out)), out)

    def test_default_docprops_fail(self):
        code, out, _ = self.run_case([S1_NOTES], [self.GOOD], author="python-pptx")
        self.assertEqual(code, 1)
        self.assertTrue(any("既定値のまま" in f for f in fails(out)), out)

    def test_slide_count_mismatch_fails(self):
        code, out, _ = self.run_case([S1_NOTES, S1], [self.GOOD])
        self.assertTrue(any("スライド数が違う" in f for f in fails(out)), out)


@unittest.skipUnless(HAS_PPTX and HAS_CHROME, "python-pptx・Node.js 22+・Chrome が要る")
class BuiltPptx(unittest.TestCase):
    """build.py で本物の pptx を書き出して読み戻す（AC-06 の機械部分）。"""

    @classmethod
    def setUpClass(cls):
        cls.td = TempDir()
        cls.d = cls.td.__enter__()
        cls.html = cls.d / "deck.html"
        cls.html.write_bytes((FIXTURES / "sample_deck.html").read_bytes())
        cls.code, cls.out, cls.err = run("build.py", cls.html, "--out", cls.d / "out")

    @classmethod
    def tearDownClass(cls):
        cls.td.__exit__(None, None, None)

    def test_build_and_readback_pass(self):
        self.assertEqual(self.code, 0, self.err)
        for name in ("deck.pptx", "deck.pdf", "deck_読み原稿.md", "manifest.json"):
            self.assertTrue((self.d / "out" / name).exists(), name)
        code, out, err = check("pptx", self.d / "out" / "deck.pptx", self.html)
        self.assertEqual(code, 0, out + err)
        self.assertIn("ノート 7 枚", out)
        self.assertIn("画像 0", out)

    def test_text_is_text_and_tables_native(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(self.d / "out" / "deck.pptx")
        n_text = sum(1 for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        n_tables = sum(1 for s in prs.slides for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table)
        n_pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertGreater(n_text, 30)
        self.assertGreaterEqual(n_tables, 1)
        self.assertEqual(n_pics, 0)
        self.assertEqual(prs.core_properties.author, "")

    def test_damaged_copy_fails(self):
        from pptx import Presentation
        prs = Presentation(self.d / "out" / "deck.pptx")
        # 3 枚目の一番長い文字の箱を消す（書き出しで文字が落ちた状態）
        s = prs.slides[2]
        boxes = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        victim = max(boxes, key=lambda sh: len(sh.text_frame.text))
        victim._element.getparent().remove(victim._element)
        bad = self.d / "bad.pptx"
        prs.save(bad)
        code, out, _ = check("pptx", bad, self.html)
        self.assertEqual(code, 1)
        self.assertTrue(any(f.startswith("FAIL p3: HTML の文字が pptx に無い") for f in fails(out)), out)


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(HAS_PPTX and HAS_CHROME, "python-pptx・Node.js 22+・Chrome が要る")
class ZoomedParts(unittest.TestCase):
    """U3 報告の回帰: 追加パーツ集（.slide{zoom:0.8}）の文字が 1.25 倍で書き出され、はみ出しで落ちていた。"""

    def test_more_parts_build_and_check(self):
        with TempDir() as d:
            code, out, err = run("new_deck.py", "--parts", "m01,m12,m21,m35,b06", "--title", "試験", "-o", d / "m.html")
            self.assertEqual(code, 0, err)
            code, out, err = run("build.py", d / "m.html", "--out", d / "out", "--formats", "pptx")
            self.assertEqual(code, 0, err)
            code, out, err = check("pptx", d / "out" / "m.pptx", d / "m.html")
            self.assertEqual(code, 0, out + err)
            from pptx import Presentation
            prs = Presentation(d / "out" / "m.pptx")
            title = [sh for sh in prs.slides[0].shapes if sh.name == "Title"][0]
            size = title.text_frame.paragraphs[0].runs[0].font.size.pt
            self.assertLessEqual(size, 23.0)   # zoom 0.8 を掛けた大きさ（掛けないと 27.5pt）
