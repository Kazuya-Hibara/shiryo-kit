"""Rebuild the binary fixtures under tests/fixtures/past-materials/.
All content is fictional. Usage: python3 build_past_materials.py <out_dir>
Needs python-pptx, python-docx, and (for the PDF) LibreOffice `soffice` on PATH."""
import subprocess, sys, tempfile, shutil, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Pt as DPt
from docx.oxml.ns import qn

OUT = pathlib.Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
FONT = "Yu Gothic"
FOOT = "※2025年9月時点の情報です。最新は必ず厚生労働省の資料でご確認ください。"
CLOSE = "迷ったら、まず相談。小さいうちなら、だいたい安く済みます。"

# ---------------------------------------------------------------- pptx
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def tb(slide, x, y, w, h, text, size=20, bold=False, color=(0x22, 0x22, 0x22)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(8)
    return box

def slide(title, body, notes, foot=True, size=22, table=None):
    s = prs.slides.add_slide(blank)
    tb(s, 0.7, 0.4, 11.9, 1.2, title, size=32, bold=True, color=(0x1F, 0x3A, 0x5F))
    if body:
        tb(s, 0.7, 1.8, 11.9, 3.2 if table else 4.6, body, size=size)
    if table:
        rows, cols = len(table), len(table[0])
        gt = s.shapes.add_table(rows, cols, Inches(0.7), Inches(3.6), Inches(11.9), Inches(0.6 * rows))
        for r in range(rows):
            for c in range(cols):
                cell = gt.table.cell(r, c)
                cell.text = table[r][c]
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(18); run.font.name = FONT
    if foot:
        tb(s, 0.7, 6.8, 11.9, 0.5, FOOT, size=12, color=(0x66, 0x66, 0x66))
    s.notes_slide.notes_text_frame.text = notes
    return s

slide("36協定、出していれば安心? ― 2025年版・中小企業の時間外労働ルール",
      ["講師: 特定社会保険労務士 柏木 遼(つきがおか労務オフィス)", "2025年9月 商工会 労務セミナー"],
      "皆さん、こんにちは。社労士の柏木です。今日は90分、36協定の話をします。\n"
      "最初にひとつだけ質問させてください。36協定、最後に見たのはいつですか? 手を挙げなくて大丈夫です。心の中でどうぞ。",
      foot=False)
slide("Q. 御社の36協定、最後に見たのはいつですか?",
      ["突然ですが、届出は出しっぱなしになっていませんか?",
       "ざっくり言うと、届出は「免許」、毎月の運用は「運転」です。免許があっても、運転が荒ければ捕まります。"],
      "ここは笑いを取るところ。免許と運転のたとえ話。家計簿でも言い換えられる(届出=予算、実績=家計簿)。\n"
      "会場の前列にいらした 株式会社サンプル商事 様 の総務 山田 花子 様 に、先月の相談の話を振る(了承済み)。")
slide("ポイントは3つ ― 上限・特別条項・記録",
      ["1. 原則の上限は 月45時間・年360時間",
       "2. 特別条項があっても、年720時間・単月100時間未満(休日労働含む)・2〜6か月平均80時間以内・月45時間超は年6回まで",
       "3. 守れているかは「記録」で決まる(タイムカード・PCログ)"],
      "ポイントは3つです、と必ず言う。指を3本立てる。上限の数字の直前に「ここ、テストに出ます」と言う。数字は黒板(ホワイトボード)に書き出して、あとで消さない。\n"
      "出典: 厚生労働省のリーフレット「時間外労働の上限規制 わかりやすい解説」(スライド右下に毎回出す。ここでは出典URLの記載なし→過去資料の弱点)。",
      size=22)
slide("原則と特別条項、どこが違う?",
      None,
      "表で見せる。特別条項は『臨時的な特別の事情』が必要、と一言添える。",
      table=[["", "原則", "特別条項つき"],
             ["月", "45時間まで", "100時間未満(休日労働含む)"],
             ["年", "360時間まで", "720時間まで"],
             ["回数", "―", "月45時間超は年6回まで"]])
slide("【落とし穴】特別条項があれば、いつでも月80時間OK? 〇か✕か",
      ["答えは ✕ です。", "2〜6か月の平均が80時間以内、という縛りが別にあります。",
       "「ひと月ならセーフ」で安心すると、平均でアウトになります。"],
      "〇✕クイズ。会場に挙手させる。8割が〇に挙げるのでそこで『実は✕です』と間を置く。\n"
      "相談事例: サンプル商事 様(山田 花子 様 / 携帯 090-0000-1234 / hanako.yamada@example.com / 〒000-0002 例県サンプル市見本町4-5-6 従業員28名)は繁忙期に月60時間超が3か月続いた。")
slide("今日持ち帰るのは、1つだけ",
      ["自社の36協定を開いて、「有効期間」と「上限の数字」に付箋を貼る。",
       "ざっくり言うと、それだけで来月の自分がラクになります。"],
      "『全部やってください』とは絶対に言わない。持ち帰りは1つだけ、という約束を守る。")
slide("まとめ", [CLOSE], "最後の一言は毎回これ。声のトーンを落として、ゆっくり。笑顔で締める。")
slide("ご相談はお気軽に",
      ["つきがおか労務オフィス  代表 柏木 遼",
       "〒000-0001 例県サンプル市見本町1-2-3 サンプルビル2F",
       "TEL 050-0000-0100 / 携帯 090-0000-0001",
       "info@tsukigaoka-romu.example.com"],
      "事務所の連絡先(事務所自身の情報)。最後に名刺を配る。", foot=False)
prs.core_properties.author = "つきがおか労務オフィス(架空)"
prs.core_properties.title = "36協定セミナー資料 2025年版(架空サンプル)"
prs.core_properties.comments = "FICTIONAL fixture for shiryo-kit tests. All names, numbers, addresses are invented."
prs.save(OUT / "2025-09_36kyotei-seminar.pptx")

# ---------------------------------------------------------------- docx newsletter
def mkdoc():
    d = Document()
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = DPt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return d
nl = mkdoc()
nl.core_properties.author = "つきがおか労務オフィス(架空)"
nl.core_properties.comments = "FICTIONAL fixture for shiryo-kit tests."
nl.add_heading("つきがおか通信 2026年3月号", 0)
nl.add_paragraph("発行: つきがおか労務オフィス  代表 特定社会保険労務士 柏木 遼")
nl.add_heading("年度末、有給休暇の「5日」は大丈夫ですか?", 1)
nl.add_paragraph("年10日以上の有給休暇が付与される従業員には、年5日の取得が会社の義務です。ご存じでしたか? "
                 "ざっくり言うと、「本人任せにせず、会社が日を決めて取ってもらう」のが基本です。")
nl.add_paragraph("ポイントは3つです。")
for t in ["対象は、年10日以上の有給が付与される従業員(パートも含みます)",
          "付与日から1年以内に5日。管理簿を作って3年保存",
          "時季指定をするときは、本人の意見を聞く"]:
    nl.add_paragraph(t, style="List Number")
nl.add_paragraph("【落とし穴】「本人が取りたがらないから仕方ない」は通りません。")
nl.add_paragraph("今月持ち帰るのは1つだけ。管理簿の「5日に届いていない人」に付箋を貼ってみてください。")
nl.add_heading("お客様の声", 1)
nl.add_paragraph("ミナト運輸 株式会社 様 総務担当 佐藤 健 様より:「『ここ、テストに出ます』と言われた上限の話、"
                 "社内で回しました。助かりました」")
nl.add_paragraph("サンプル食品 株式会社 様 田中 美咲 様(03-0000-0456 / misaki.tanaka@example.com)には、"
                 "来月の研修でお話しいただく予定です。")
nl.add_heading("事務所からのお知らせ", 1)
nl.add_paragraph("春の相談会 3月26日(木)14時〜。会場: 〒000-0001 例県サンプル市見本町1-2-3 サンプルビル2F")
nl.add_paragraph("お問い合わせ: 050-0000-0100 / 携帯 090-0000-0001 / info@tsukigaoka-romu.example.com")
nl.add_paragraph("迷ったら、まず相談。小さいうちなら、だいたい安く済みます。")
nl.add_paragraph("※2026年3月時点の情報です。最新は必ず厚生労働省の資料でご確認ください。")
nl.save(OUT / "2026-03_newsletter.docx")

# ---------------------------------------------------------------- pdf checksheet (docx -> pdf)
cs = mkdoc()
cs.add_heading("36協定 5分チェックシート", 0)
cs.add_paragraph("ご記入: サンプル食品 株式会社 様  担当 田中 美咲 様  (03-0000-0456)")
cs.add_paragraph("突然ですが、次の5つに全部「はい」と言えますか?")
tbl = cs.add_table(rows=1, cols=2); tbl.style = "Table Grid"
tbl.rows[0].cells[0].text = "確認すること"; tbl.rows[0].cells[1].text = "はい / いいえ"
for q in ["36協定を労基署に届け出て、有効期間内である",
          "原則の上限(月45時間・年360時間)を超えた月がない",
          "特別条項を使った月は年6回以内である",
          "2〜6か月の平均が80時間以内に収まっている",
          "労働時間の記録(タイムカード等)を3年保存している"]:
    r = tbl.add_row().cells; r[0].text = q; r[1].text = "はい / いいえ"
cs.add_paragraph("【落とし穴】いちばん多いのは、4つ目です。ひと月ごとには収まっていても、平均で超えます。")
cs.add_paragraph("ざっくり言うと、「いいえ」が1つでもあれば、今日の持ち帰りはそれ1つだけです。")
cs.add_paragraph("迷ったら、まず相談。小さいうちなら、だいたい安く済みます。  つきがおか労務オフィス 050-0000-0100")
cs.add_paragraph("※2025年9月時点の情報です。最新は必ず厚生労働省の資料でご確認ください。")
tmp = pathlib.Path(tempfile.mkdtemp())
cs.save(tmp / "2025-10_36kyotei-checksheet.docx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp),
                str(tmp / "2025-10_36kyotei-checksheet.docx")], check=True, capture_output=True, timeout=180)
shutil.copy(tmp / "2025-10_36kyotei-checksheet.pdf", OUT / "2025-10_36kyotei-checksheet.pdf")
print("built", sorted(p.name for p in OUT.iterdir()))
